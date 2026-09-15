"""Multimedia processing via in-process models and lightweight document parsers."""

# pylint: disable=wrong-import-order,import-outside-toplevel
import os
import tempfile
import csv
from collections.abc import Callable

from app.core.config import settings
from app.core.log import get_logger

logger = get_logger("MultimediaService")


def _format_timestamp(seconds: float) -> str:
    """Convert seconds to M:SS or H:MM:SS string."""
    s = int(seconds)
    m, s = divmod(s, 60)
    h, m = divmod(m, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def image_data_url(image_path: str) -> str:
    """JPEG data URL of ``image_path``, downscaled to the describe budget.

    Every provider takes a data URL; sending a 12-megapixel photo to a local
    projector or a metered endpoint buys nothing over ~1.5 MP.
    """
    import base64
    import io

    from PIL import Image

    max_pixels = int(getattr(settings, "IMAGE_DESCRIBE_MAX_PIXELS", 0) or 1_500_000)
    with Image.open(image_path) as opened:
        image = opened.convert("RGB")
    pixels = image.width * image.height
    if max_pixels > 0 and pixels > max_pixels:
        scale = (max_pixels / pixels) ** 0.5
        image = image.resize(
            (max(1, int(image.width * scale)), max(1, int(image.height * scale))),
            Image.LANCZOS,
        )
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=88)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


class MultimediaService:
    """Extract text from attachments via in-process models and vault files."""

    def _resolve_vault_local_path(self, path_or_url: str) -> str | None:
        """Map ``/vault-files/<kb_id>/<rel>`` to an absolute vault file path."""
        from pathlib import Path
        from urllib.parse import unquote, urlparse

        if not path_or_url:
            return None

        raw = path_or_url.strip()
        if raw.startswith("http://") or raw.startswith("https://"):
            raw = urlparse(raw).path or ""

        raw = raw.split("?", 1)[0].replace("\\", "/")
        marker = "/vault-files/"
        if marker not in raw and not raw.startswith("vault-files/"):
            return None

        if raw.startswith("vault-files/"):
            raw = "/" + raw
        idx = raw.find(marker)
        remainder = raw[idx + len(marker) :] if idx >= 0 else ""
        parts = remainder.split("/", 1)
        if len(parts) < 2 or not parts[0] or not parts[1]:
            return None

        kb_id, rel = parts[0], unquote(parts[1]).lstrip("/")
        if not rel or ".." in Path(rel).parts:
            return None

        try:
            from app.services.kb_registry import kb_registry

            kb = kb_registry.get_kb(kb_id)
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.warning(f"Vault lookup failed for kb={kb_id}: {exc}")
            return None

        if not kb or not kb.vault_path:
            logger.warning(f"No vault for kb={kb_id} (url={path_or_url})")
            return None

        vault_root = Path(kb.vault_path).resolve()
        abs_path = (vault_root / rel).resolve()
        try:
            abs_path.relative_to(vault_root)
        except ValueError:
            logger.warning(f"Rejected path escape: {abs_path} not under {vault_root}")
            return None

        if abs_path.is_file():
            return str(abs_path)
        logger.warning(f"Vault file missing: {abs_path}")
        return None

    def _is_ephemeral_download(self, original_ref: str, local_path: str) -> bool:
        """True when ``local_path`` is a temp download we must delete (not a vault file)."""
        if not local_path or local_path == original_ref:
            return False
        if self._resolve_vault_local_path(original_ref) == local_path:
            return False
        if os.path.isfile(original_ref) and os.path.samefile(original_ref, local_path):
            return False
        return True

    # Remote attachments are media files; anything past this is more likely a
    # hostile/exhausting URL than a legitimate embed.
    _MAX_REMOTE_DOWNLOAD_BYTES = 512 * 1024 * 1024

    @staticmethod
    def _assert_public_http_url(url: str) -> None:
        """Reject URLs that resolve to loopback/private/link-local addresses.

        Note markdown can contain arbitrary links (including web-ingested
        content), so ingest-time fetches must not be able to reach local
        services (API, Qdrant, Meilisearch, Firefly) or the LAN.
        """
        import ipaddress
        import socket
        from urllib.parse import urlparse

        parsed = urlparse(url)
        if parsed.scheme not in ("http", "https") or not parsed.hostname:
            raise ValueError(f"Unsupported attachment URL: {url}")
        try:
            infos = socket.getaddrinfo(parsed.hostname, parsed.port or 0)
        except socket.gaierror as exc:
            raise ValueError(f"Cannot resolve attachment host: {parsed.hostname}") from exc
        for info in infos:
            addr = ipaddress.ip_address(info[4][0])
            if (
                addr.is_loopback
                or addr.is_private
                or addr.is_link_local
                or addr.is_reserved
                or addr.is_multicast
                or addr.is_unspecified
            ):
                raise ValueError(
                    f"Refusing to fetch non-public address for {parsed.hostname}: {addr}"
                )

    def _download_temp_file(self, path_or_url: str) -> str:
        """Resolve vault/local/remote attachments to a local file path."""
        import tempfile

        import requests

        if os.path.isfile(path_or_url):
            return path_or_url

        vault_path = self._resolve_vault_local_path(path_or_url)
        if vault_path:
            logger.info(f"Resolved vault attachment: {vault_path}")
            return vault_path

        if not path_or_url.startswith("http"):
            raise FileNotFoundError(f"Attachment not found: {path_or_url}")

        self._assert_public_http_url(path_or_url)
        logger.info(f"Downloading remote file: {path_or_url}...")
        suffix = "." + path_or_url.split(".")[-1] if "." in path_or_url else ".tmp"
        with requests.get(path_or_url, timeout=300, stream=True) as response:
            response.raise_for_status()
            written = 0
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                try:
                    for chunk in response.iter_content(chunk_size=1024 * 1024):
                        written += len(chunk)
                        if written > self._MAX_REMOTE_DOWNLOAD_BYTES:
                            raise ValueError(
                                f"Remote attachment exceeds "
                                f"{self._MAX_REMOTE_DOWNLOAD_BYTES} bytes: {path_or_url}"
                            )
                        tmp.write(chunk)
                except Exception:
                    tmp.close()
                    os.unlink(tmp.name)
                    raise
                return tmp.name

    def _caption_video_with_marlin(self, local_path: str) -> dict:
        """Marlin caption via in-process multimodal runtime."""
        from app.services.multimodal_runtime import multimodal_runtime

        try:
            return multimodal_runtime.caption_video_path(local_path)
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.error(f"Marlin captioning failed: {exc}")
            raise RuntimeError(f"Marlin captioning failed: {exc}") from exc

    def unload_local_models(self, family: str | None = None) -> None:
        """Unload the speech model (or every media model) from the API process."""
        from app.services.multimodal_runtime import multimodal_runtime

        try:
            multimodal_runtime.unload(family)
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.warning(f"Multimodal unload skipped/failed: {exc}")

    def unload_marlin(self) -> None:
        """Unload Marlin from the API process."""
        from app.services.multimodal_runtime import multimodal_runtime

        try:
            multimodal_runtime.unload("marlin")
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.warning(f"Marlin unload skipped/failed: {exc}")

    def describe_image(self, image_path: str, llm) -> str:
        """Describe an image with the KB's ingestion model (see ``LLMService.describe_image``)."""
        local_path = self._download_temp_file(image_path)
        try:
            text = llm.describe_image(local_path)
            if not text:
                raise RuntimeError("The model returned no description")
            return text
        finally:
            if self._is_ephemeral_download(image_path, local_path) and os.path.exists(
                local_path
            ):
                os.remove(local_path)

    def transcribe_audio(self, audio_path: str) -> str:
        """Transcribe audio via in-process Qwen3-ASR (multimodal_runtime)."""
        local_path = self._download_temp_file(audio_path)
        try:
            logger.info(f"Transcribing audio: {local_path}")
            from app.services.multimodal_runtime import multimodal_runtime

            return multimodal_runtime.transcribe_audio_path(local_path) or ""
        finally:
            if self._is_ephemeral_download(audio_path, local_path) and os.path.exists(
                local_path
            ):
                os.remove(local_path)

    def process_video(self, video_path: str) -> str:
        """Process a video with Qwen3-ASR audio transcription and Marlin visual analysis."""
        transcript = self.transcribe_video_audio(video_path)
        visual = self.describe_video_visual(video_path)
        parts = []
        if transcript:
            parts.append(f"### Spoken Content\n{transcript}")
        if visual:
            parts.append(visual)
        return "\n\n".join(parts) if parts else "(Video processing produced no output)"

    def transcribe_video_audio(self, video_path: str) -> str:
        """Transcribe a video's audio track without running visual analysis."""
        import av

        local_path = self._download_temp_file(video_path)
        try:
            with av.open(local_path) as container:
                has_video = len(container.streams.video) > 0

            if not has_video:
                logger.info("No video stream detected — treating as audio-only.")
                return self.transcribe_audio(local_path)

            try:
                logger.info("Transcribing video audio track...")
                return self.transcribe_audio(local_path)
            except Exception as exc:  # pylint: disable=broad-exception-caught
                logger.warning(
                    f"Audio transcription skipped (no audio track or failed): {exc}"
                )
                return ""
        finally:
            if self._is_ephemeral_download(video_path, local_path) and os.path.exists(
                local_path
            ):
                os.remove(local_path)

    def describe_video_visual(self, video_path: str) -> str:
        """Run Marlin visual analysis without transcribing audio."""
        import av

        local_path = self._download_temp_file(video_path)
        try:
            with av.open(local_path) as container:
                has_video = len(container.streams.video) > 0

            if not has_video:
                return ""

            try:
                logger.info("Running Marlin video captioning via service...")
                result = self._caption_video_with_marlin(local_path)
                scene = result.get("scene", "")
                events = result.get("events", [])
                lines = [
                    f"- {_format_timestamp(ev.get('start', 0))}\u2013{_format_timestamp(ev.get('end', 0))} \u2014 {ev.get('description', '')}"
                    for ev in events
                ]
                events_text = "\n".join(lines)
                logger.info(f"Marlin: {len(events)} events extracted.")
            except Exception as exc:  # pylint: disable=broad-exception-caught
                logger.error(f"Marlin captioning failed: {exc}")
                return ""

            if scene:
                visual = f"### Visual Analysis\n**Scene:** {scene}"
                if events_text:
                    visual += f"\n\n**Events:**\n{events_text}"
                return visual
            return ""
        finally:
            if self._is_ephemeral_download(video_path, local_path) and os.path.exists(
                local_path
            ):
                os.remove(local_path)

    def _pdf_page_needs_render(self, page, native_text: str, image_descriptions: list[str]) -> bool:
        """True when describing a full-page render is useful (scanned / sparse pages)."""
        if not settings.PDF_VISUAL_EXTRACTION_ENABLED:
            return False
        if len(native_text.strip()) >= settings.PDF_VISUAL_TEXT_THRESHOLD:
            return False
        # Embedded-image descriptions already covered this page.
        if image_descriptions:
            return False
        try:
            if page.get_images(full=True):
                return True
        except Exception:  # pylint: disable=broad-exception-caught
            pass
        try:
            if page.get_drawings():
                return True
        except Exception:  # pylint: disable=broad-exception-caught
            pass
        # Image-only / empty pages with almost no text still benefit from a render.
        return len(native_text.strip()) == 0

    def _describe_pdf_page_render(self, page, llm) -> str:
        """Render a PDF page to PNG and describe it with the ingestion model."""
        import tempfile

        import fitz

        dpi = max(int(settings.PDF_VISUAL_RENDER_DPI or 144), 72)
        zoom = dpi / 72.0
        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(zoom, zoom),
            alpha=False,
            annots=True,
        )
        image_path = ""
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
                image_path = tmp.name
                pixmap.save(image_path)
            return (llm.describe_image(image_path) or "").strip()
        finally:
            if image_path and os.path.exists(image_path):
                os.remove(image_path)

    def extract_text_from_pdf(
        self,
        pdf_path: str,
        progress_callback: Callable[[str, str | None], None] | None = None,
        llm=None,
    ) -> str:
        """Extract PDF page text; the ingestion model reads embedded images and sparse page renders."""
        import tempfile

        import fitz

        def _progress(stage: str, model: str | None = None) -> None:
            if progress_callback:
                progress_callback(stage, model)

        local_path = self._download_temp_file(pdf_path)
        owns_temp = self._is_ephemeral_download(pdf_path, local_path)
        vision_model = (llm.get_ingestion_model() if llm else None) or "vision model"
        try:
            extracted_pages: list[str] = []
            doc = fitz.open(local_path)
            total_pages = len(doc)
            max_visual_pages = int(settings.PDF_VISUAL_EXTRACTION_MAX_PAGES or 0)
            visual_pages_used = 0
            try:
                for page_index, page in enumerate(doc, start=1):
                    _progress(
                        f"PDF: page {page_index}/{total_pages}, extracting text",
                        None,
                    )
                    native_text = page.get_text().strip()
                    page_parts = [f"--- Page {page_index} ---"]
                    image_descriptions: list[str] = []

                    if native_text:
                        page_parts.append(f"Native text:\n{native_text}")

                    images = page.get_images(full=True)
                    for image_index, image_info in enumerate(images, start=1):
                        _progress(
                            (
                                f"PDF: page {page_index}/{total_pages}, "
                                f"describing image {image_index}/{len(images)}"
                            ),
                            vision_model,
                        )
                        xref = image_info[0]
                        image_path = ""
                        try:
                            extracted = doc.extract_image(xref)
                            image_bytes = extracted.get("image")
                            if not image_bytes:
                                continue
                            image_ext = extracted.get("ext") or "png"
                            with tempfile.NamedTemporaryFile(
                                delete=False, suffix=f".{image_ext}"
                            ) as tmp:
                                image_path = tmp.name
                                tmp.write(image_bytes)
                            description = llm.describe_image(image_path) if llm else ""
                        except Exception as exc:  # pylint: disable=broad-exception-caught
                            logger.warning(
                                "PDF image description skipped "
                                f"(page={page_index}, image={image_index}): {exc}"
                            )
                            continue
                        finally:
                            if image_path and os.path.exists(image_path):
                                os.remove(image_path)

                        if description:
                            image_descriptions.append(
                                f"Image {image_index}: {description}"
                            )

                    # Scanned / sparse pages: render whole page when embeds yielded nothing.
                    if self._pdf_page_needs_render(page, native_text, image_descriptions):
                        if not max_visual_pages or visual_pages_used < max_visual_pages:
                            _progress(
                                f"PDF: page {page_index}/{total_pages}, describing page render",
                                vision_model,
                            )
                            try:
                                page_desc = self._describe_pdf_page_render(page, llm) if llm else ""
                                if page_desc:
                                    image_descriptions.append(
                                        f"Page render: {page_desc}"
                                    )
                                    visual_pages_used += 1
                            except Exception as exc:  # pylint: disable=broad-exception-caught
                                logger.warning(
                                    f"PDF page render description skipped "
                                    f"(page={page_index}): {exc}"
                                )

                    if image_descriptions:
                        page_parts.append(
                            "Image descriptions:\n" + "\n".join(image_descriptions)
                        )

                    if native_text or image_descriptions:
                        extracted_pages.append("\n\n".join(page_parts))

                    _progress(f"PDF: page {page_index}/{total_pages} complete", None)
            finally:
                doc.close()

            full_text = "\n\n".join(extracted_pages).strip()
            if not full_text:
                return "PDF contains no extractable native or visual content."
            return full_text
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.error(f"PDF extraction failed: {exc}")
            raise RuntimeError(f"PDF extraction failed: {exc}") from exc
        finally:
            if owns_temp and local_path and os.path.exists(local_path):
                os.remove(local_path)

    def extract_text_from_docx(self, docx_path: str) -> str:
        """Extract text from a Word document (.docx) using native parsing only.

        Covers the document body, tables, headers/footers and text boxes.
        ``python-docx`` walks only ``document.paragraphs`` by default, so
        anything in a header, footer or floating text box is invisible to it —
        which for many real documents is where the title, author and captions
        live.
        """
        local_path = self._download_temp_file(docx_path)
        parts = []

        try:
            try:
                import docx

                document = docx.Document(local_path)
                for paragraph in document.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)

                for idx, table in enumerate(document.tables, start=1):
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            rows.append(" | ".join(cells))
                    if rows:
                        parts.append(f"--- Table {idx} ---\n" + "\n".join(rows))

                # Headers and footers, deduped: repeating the same running head
                # on every section adds nothing but tokens.
                seen_chrome: set[str] = set()
                for label, attr in (("Header", "header"), ("Footer", "footer")):
                    lines = []
                    for section in document.sections:
                        for para in getattr(section, attr).paragraphs:
                            text = para.text.strip()
                            if text and text not in seen_chrome:
                                seen_chrome.add(text)
                                lines.append(text)
                    if lines:
                        parts.append(f"--- {label} ---\n" + "\n".join(lines))

                box_text = self._docx_text_boxes(document)
                if box_text:
                    parts.append("--- Text boxes ---\n" + "\n".join(box_text))

                full_text = "\n\n".join(parts).strip()
                if not full_text:
                    return "Word document contains no extractable text."
                return full_text

            except ImportError as exc:
                raise RuntimeError(
                    "Word extraction unavailable: python-docx is not installed."
                ) from exc
            except Exception as exc:  # pylint: disable=broad-exception-caught
                logger.error(f"DOCX extraction failed: {exc}")
                raise RuntimeError(f"Word extraction failed: {exc}") from exc

        finally:
            if self._is_ephemeral_download(docx_path, local_path) and os.path.exists(
                local_path
            ):
                try:
                    os.remove(local_path)
                except OSError:
                    pass

    @staticmethod
    def _docx_text_boxes(document) -> list[str]:
        """Text inside floating shapes, which python-docx has no API for."""
        ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        seen: set[str] = set()
        out: list[str] = []
        try:
            for node in document.element.body.findall(".//w:txbxContent", ns):
                runs = [t.text or "" for t in node.findall(".//w:t", ns)]
                text = "".join(runs).strip()
                if text and text not in seen:
                    seen.add(text)
                    out.append(text)
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.warning(f"DOCX text-box scan failed: {exc}")
        return out

    def extract_docx_images(self, docx_path: str, max_images: int = 20) -> list[str]:
        """Write a .docx's embedded images to temp files and return their paths.

        A diagram inside a Word document used to be invisible to the pipeline
        while the same file attached directly got described. The
        caller is responsible for deleting the returned paths.
        """
        local_path = self._download_temp_file(docx_path)
        written: list[str] = []
        try:
            import docx

            document = docx.Document(local_path)
            for name, part in document.part.related_parts.items():
                if len(written) >= max_images:
                    logger.info(
                        "DOCX image cap reached (%d) — skipping the rest", max_images
                    )
                    break
                if "image" not in getattr(part, "content_type", ""):
                    continue
                blob = getattr(part, "blob", None)
                # Skip spacers and bullets; they describe as noise.
                if not blob or len(blob) < 8192:
                    continue
                ext = os.path.splitext(str(getattr(part, "partname", name)))[1] or ".png"
                fd, tmp = tempfile.mkstemp(suffix=ext, prefix="orb_docx_img_")
                with os.fdopen(fd, "wb") as handle:
                    handle.write(blob)
                written.append(tmp)
        except Exception as exc:  # pylint: disable=broad-exception-caught
            logger.warning(f"DOCX image extraction failed: {exc}")
        finally:
            if self._is_ephemeral_download(docx_path, local_path) and os.path.exists(
                local_path
            ):
                try:
                    os.remove(local_path)
                except OSError:
                    pass
        return written

    def extract_text_from_spreadsheet(
        self, sheet_path: str
    ) -> str:  # pylint: disable=too-many-return-statements,too-many-nested-blocks,too-many-locals,too-many-branches
        """Extract text from spreadsheet-like files using native parsing."""
        local_path = self._download_temp_file(sheet_path)
        lower_path = local_path.lower()

        try:  # pylint: disable=too-many-nested-blocks
            if lower_path.endswith(".xlsx"):
                try:
                    from openpyxl import load_workbook

                    workbook = load_workbook(
                        filename=local_path,
                        read_only=True,
                        data_only=True,
                    )
                    parts = []
                    for sheet_name in workbook.sheetnames:
                        sheet = workbook[sheet_name]
                        rows = []
                        for row in sheet.iter_rows(values_only=True):
                            values = [
                                str(cell).strip() for cell in row if cell is not None
                            ]
                            if values:
                                rows.append("\t".join(values))
                        if rows:
                            parts.append(
                                f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows)
                            )

                    workbook.close()
                    full_text = "\n\n".join(parts).strip()
                    if not full_text:
                        return "Spreadsheet contains no extractable text."
                    return full_text

                except ImportError as exc:
                    raise RuntimeError(
                        "Spreadsheet extraction unavailable: openpyxl is not installed."
                    ) from exc
                except Exception as exc:  # pylint: disable=broad-exception-caught
                    logger.error(f"XLSX extraction failed: {exc}")
                    raise RuntimeError(f"Spreadsheet extraction failed: {exc}") from exc

            if lower_path.endswith((".csv", ".tsv")):
                delimiter = "\t" if lower_path.endswith(".tsv") else ","
                rows = []
                with open(local_path, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.reader(f, delimiter=delimiter)
                    for row in reader:
                        values = [cell.strip() for cell in row if cell and cell.strip()]
                        if values:
                            rows.append("\t".join(values))

                if not rows:
                    return "Spreadsheet contains no extractable text."
                return "\n".join(rows)

            if lower_path.endswith(".xls"):
                raise RuntimeError(
                    "Legacy .xls is not supported. Please convert to .xlsx."
                )

            raise RuntimeError("Unsupported spreadsheet format.")
        finally:
            if self._is_ephemeral_download(sheet_path, local_path) and os.path.exists(
                local_path
            ):
                os.remove(local_path)


multimedia_service = MultimediaService()
